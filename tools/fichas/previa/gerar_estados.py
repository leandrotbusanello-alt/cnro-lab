#!/usr/bin/env python3
"""
Dados de teste de uma ficha a partir dos valores digitados na própria planilha (quando ela tem exemplo).
Uso (na raiz do repositório): python3 tools/fichas/previa/gerar_estados.py FR-IMOB-19_Rev00 [...]
Grava tools/fichas/previa/estados/<ficha>.json. Para fichas sem exemplo, escreva o JSON à mão:
  { "entradas": { "D18": 65.13, "VERSO!C6": "BAL-01" }, "escolhas": { "grupo": "Sim" }, "verificacoes": { "VERSO!A9": true } }
"""
import datetime, json, sys
from pathlib import Path
import openpyxl

AQUI = Path(__file__).resolve().parent
FICHAS = AQUI.parent
for f in sys.argv[1:]:
    cod = f.rsplit('_', 1)[0]
    spec = json.loads((FICHAS / 'specs' / f'{cod}.json').read_text(encoding='utf-8'))
    m = json.loads((FICHAS / 'saida' / f'{f}.modelo.json').read_text(encoding='utf-8'))['modelo']
    wb = openpyxl.load_workbook(FICHAS / 'planilhas' / spec['arquivo'])
    abas = {'': spec.get('aba') or wb.sheetnames[0]}
    for ab in spec.get('abas_extras', []):
        abas[ab['id']] = ab['aba']
    e = {}
    for fid, fo in [('', m)] + [(ab['id'], ab) for ab in m.get('abas', [])]:
        ws = wb[abas[fid]]
        pre = f'{fid}!' if fid else ''
        for a, d in fo['cells'].items():
            if (d.get('role') or {}).get('tipo') not in ('entrada', 'revisao'):
                continue
            v = ws[a].value
            if v is None or (isinstance(v, str) and v.startswith('=')):
                continue
            if isinstance(v, datetime.datetime):
                v = (v - datetime.datetime(1899, 12, 30)).total_seconds() / 86400
            elif isinstance(v, datetime.time):
                v = (v.hour * 3600 + v.minute * 60 + v.second) / 86400
            e[pre + a] = v
    (AQUI / 'estados' / f'{f}.json').write_text(json.dumps({'entradas': e}, ensure_ascii=False), encoding='utf-8')
    print(f, len(e), 'entradas')
