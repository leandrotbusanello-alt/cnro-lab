#!/usr/bin/env python3
"""
Conversor de fichas — Excel original (.xlsx) → modelo da ficha online (JSON).

Uso:
    python3 tools/fichas/converter.py                 # converte todas as specs
    python3 tools/fichas/converter.py FR-IMOB-13      # só uma

Para cada tools/fichas/specs/<CODIGO>.json gera:
    tools/fichas/saida/<CODIGO>_<VERSAO>.modelo.json       modelo publicado no banco
    tools/fichas/saida/<CODIGO>_<VERSAO>.verificacao.json  dados do exemplo + valores do Excel (não vai ao banco)
e, ao final, supabase/migrations/13b_fichas_modelo_carga.sql com todos os modelos.

O Excel é a fonte da verdade do layout e das fórmulas: nada é redesenhado.
A spec só diz o papel de cada célula (quem preenche o quê) e como gerar os resultados.
Requisitos: Python 3.10+, openpyxl, lxml.
"""
import base64, datetime, hashlib, json, re, sys
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries
from lxml import etree

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
MOTOR = 1          # versão do formato do modelo / motor de fichas
EMU_PX = 9525
MDW = 7            # largura máxima de dígito (px) da fonte padrão

INDEXED = ['000000','FFFFFF','FF0000','00FF00','0000FF','FFFF00','FF00FF','00FFFF','000000','FFFFFF','FF0000','00FF00','0000FF','FFFF00','FF00FF','00FFFF','800000','008000','000080','808000','800080','008080','C0C0C0','808080','9999FF','993366','FFFFCC','CCFFFF','660066','FF8080','0066CC','CCCCFF','000080','FF00FF','FFFF00','00FFFF','800080','800000','008080','0000FF','00CCFF','CCFFFF','CCFFCC','FFFF99','99CCFF','FF99CC','CC99FF','FFCC99','3366FF','33CCCC','99CC00','FFCC00','FF9900','FF6600','666699','969696','003366','339966','003300','333300','993300','993366','333399','333333']
FUNCOES_BLOQUEADAS = ('RANDBETWEEN(', 'RAND(', 'NOW(', 'TODAY(', 'INDIRECT(', 'OFFSET(')


# ── cores ────────────────────────────────────────────────────────────────────

def cores_do_tema(wb):
    try:
        root = etree.fromstring(wb.loaded_theme)
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        out = {}
        for el in root.find('.//a:clrScheme', ns):
            filho = el[0]
            out[etree.QName(el).localname] = filho.get('lastClr') or filho.get('val')
        ordem = ['lt1','dk1','lt2','dk2','accent1','accent2','accent3','accent4','accent5','accent6','hlink','folHlink']
        return [out.get(k, '000000') for k in ordem]
    except Exception:
        return ['FFFFFF','000000','EEECE1','1F497D','4F81BD','C0504D','9BBB59','8064A2','4BACC6','F79646','0000FF','800080']


def aplicar_tint(hexc, tint):
    if not tint:
        return hexc
    r, g, b = (int(hexc[i:i+2], 16) for i in (0, 2, 4))
    f = (lambda v: v * (1 + tint)) if tint < 0 else (lambda v: v + (255 - v) * tint)
    return ''.join('%02X' % max(0, min(255, round(f(v)))) for v in (r, g, b))


class Cores:
    def __init__(self, wb):
        self.tema = cores_do_tema(wb)

    def __call__(self, c, padrao=None):
        if c is None:
            return padrao
        try:
            if c.type == 'rgb' and isinstance(c.rgb, str):
                return '#' + c.rgb[-6:]
            if c.type == 'theme':
                return '#' + aplicar_tint(self.tema[c.theme], c.tint)
            if c.type == 'indexed':
                return padrao if c.indexed in (64, 65) else '#' + INDEXED[c.indexed]
        except Exception:
            pass
        return padrao


# ── utilitários ──────────────────────────────────────────────────────────────

def serial_excel(v):
    if isinstance(v, datetime.datetime):
        d = v - datetime.datetime(1899, 12, 30)
        return d.days + d.seconds / 86400
    if isinstance(v, datetime.date):
        return (v - datetime.date(1899, 12, 30)).days
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    return v


def endereco(c, r):
    return f'{get_column_letter(c)}{r}'


def separar(addr):
    m = re.match(r'^([A-Z]+)(\d+)$', addr)
    return column_index_from_string(m.group(1)), int(m.group(2))


def expandir_linhas(v):
    """[18, 19] | "18-21" | ["18-21", 23] → lista de linhas"""
    if isinstance(v, int):
        return [v]
    if isinstance(v, str):
        if '-' in v:
            a, b = v.split('-')
            return list(range(int(a), int(b) + 1))
        return [int(v)]
    out = []
    for x in v:
        out += expandir_linhas(x)
    return out


def celulas_do_grupo(g):
    if 'celulas' in g:
        return list(g['celulas'])
    return [f'{c}{r}' for c in g['colunas'] for r in expandir_linhas(g['linhas'])]


def larguras(ws, max_col):
    base = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.71
    w = {i: base for i in range(1, max_col + 1)}
    for d in ws.column_dimensions.values():
        for i in range(max(1, d.min or 0), min(d.max or 0, max_col) + 1):
            w[i] = 0 if d.hidden else (d.width if d.width is not None else base)
    return [0 if w[i] == 0 else int(((256 * w[i] + int(128 / MDW)) / 256) * MDW) for i in range(1, max_col + 1)]


def alturas(ws, max_row):
    base = ws.sheet_format.defaultRowHeight or 15
    out = []
    for r in range(1, max_row + 1):
        d = ws.row_dimensions[r] if r in ws.row_dimensions else None
        if d is not None and d.hidden:
            out.append(0)
        elif d is not None and d.ht is not None:
            out.append(round(d.ht * 96 / 72, 2))
        else:
            out.append(round(base * 96 / 72, 2))
    return out


def fonte(cores, f):
    d = {}
    if f is None:
        return d
    if f.name: d['n'] = f.name
    if f.sz: d['s'] = float(f.sz)
    if f.b: d['b'] = 1
    if f.i: d['i'] = 1
    if f.u: d['u'] = 1
    cor = cores(f.color)
    if cor and cor.upper() != '#000000': d['c'] = cor
    return d


def trechos(v, cores):
    out = []
    for p in v:
        if isinstance(p, TextBlock):
            d = {'t': p.text}
            f = p.font
            if f is not None:
                if f.b: d['b'] = 1
                if f.i: d['i'] = 1
                if f.vertAlign: d['va'] = f.vertAlign
                if f.sz: d['s'] = float(f.sz)
                cor = cores(f.color)
                if cor and cor.upper() != '#000000': d['c'] = cor
            out.append(d)
        else:
            out.append({'t': str(p)})
    return out


def borda(cores, s):
    if s is None or s.style is None:
        return None
    return [s.style, cores(s.color, '#000000') or '#000000']


# ── papéis das células (a partir da spec) ────────────────────────────────────

def papeis_da_spec(spec):
    p = {}
    for addr, campo in spec.get('pedido', {}).items():
        # "C7": "os"  ou  "B7": {"campo": "os", "manter_texto": true}  (rótulo e valor na mesma célula)
        if isinstance(campo, dict):
            p[addr] = {'tipo': 'pedido', 'campo': campo['campo']}
            if campo.get('manter_texto'):
                p[addr]['_manter'] = True
        else:
            p[addr] = {'tipo': 'pedido', 'campo': campo}
    for g in spec.get('entradas', []):
        for a in celulas_do_grupo(g):
            p[a] = {'tipo': 'entrada', 'dado': g.get('tipo', 'numero')}
            if g.get('multilinha'):
                p[a]['ml'] = 1
            if g.get('manter_texto'):
                p[a]['_manter'] = True
    for g in spec.get('revisao', []):
        for a in celulas_do_grupo(g):
            p[a] = {'tipo': 'revisao', 'dado': g.get('tipo', 'texto')}
    for g in spec.get('escolhas', []):
        for a, opcao in g['celulas'].items():
            p[a] = {'tipo': 'escolha', 'grupo': g['grupo'], 'opcao': opcao, 'texto': opcao}
            if g.get('marca'):          # a célula mostra a marca ("X") quando a opção está escolhida
                p[a]['marca'] = g['marca']
            if g.get('rotulo'):         # nome do grupo na visão em lista (senão, o texto à esquerda)
                p[a]['rotuloGrupo'] = g['rotulo']
    for quem, a in spec.get('assinaturas', {}).items():
        p[a] = {'tipo': 'assinatura', 'quem': quem}
    for a in spec.get('_verificacoes', []):
        # caixa de seleção independente (pontos de verificação): ☐/☒, cada uma por si
        p[a] = {'tipo': 'verificacao'}
    return p


def caixas_de_selecao(caminho, ws):
    """Caixas de seleção (controles de formulário do Excel) da aba: [(endereço da célula, marcada)].
    O controle fica num desenho VML ligado à aba; a posição vem de <x:Anchor> (coluna, dx, linha, dy, …)."""
    import zipfile, posixpath
    ns_rel = '{http://schemas.openxmlformats.org/package/2006/relationships}'
    with zipfile.ZipFile(caminho) as z:
        wbx = etree.fromstring(z.read('xl/workbook.xml'))
        rels = etree.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        alvo_rid = {r.get('Id'): r.get('Target') for r in rels.iter(ns_rel + 'Relationship')}
        rid = None
        for sh in wbx.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet'):
            if sh.get('name') == ws.title:
                rid = sh.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        if not rid:
            return []
        folha = posixpath.normpath(posixpath.join('xl', alvo_rid[rid].lstrip('/').replace('xl/', '', 1)))
        arq_rels = posixpath.join(posixpath.dirname(folha), '_rels', posixpath.basename(folha) + '.rels')
        if arq_rels not in z.namelist():
            return []
        out = []
        for r in etree.fromstring(z.read(arq_rels)).iter(ns_rel + 'Relationship'):
            if not r.get('Type', '').endswith('/vmlDrawing'):
                continue
            vml = posixpath.normpath(posixpath.join(posixpath.dirname(folha), r.get('Target')))
            texto = z.read(vml).decode('utf-8', errors='replace')
            for forma in re.findall(r'<v:shape\b.*?</v:shape>', texto, re.S):
                if not re.search(r'ObjectType="Checkbox"', forma, re.I):
                    continue
                m = re.search(r'<x:Anchor>\s*([^<]+)</x:Anchor>', forma)
                if not m:
                    continue
                col0, _, lin0 = (int(x) for x in m.group(1).split(',')[:3])
                marcada = bool(re.search(r'<x:Checked>\s*1', forma))
                out.append((endereco(col0 + 1, lin0 + 1), marcada))
        return out


# ── conversão ────────────────────────────────────────────────────────────────

def converter_folha(caminho, wb, wbv, ws, spec, cores):
    """Converte uma aba da planilha (frente, verso…). spec = papéis das células desta aba."""
    wsv = wbv[ws.title]

    # área da ficha: a área de impressão do Excel, ou "area_impressao" da spec quando a do Excel
    # inclui colunas de rascunho (p.ex. células com #REF! à direita da ficha)
    pa = spec.get('area_impressao') or ws.print_area
    pa = (pa if isinstance(pa, str) else pa[0]).split('!')[-1].replace('$', '')
    c1, r1, c2, r2 = range_boundaries(pa)
    c2_impressao = c2
    if spec.get('colunas_tela'):
        # colunas fora da área de impressão que o assistente usa (aparecem só na tela)
        #   "colunas_tela": "V:X"  → a grade vai até a coluna X
        c2 = max(c2, column_index_from_string(spec['colunas_tela'].split(':')[-1].strip()))
    colpx = larguras(ws, c2)[c1 - 1:]
    rowpx = alturas(ws, r2)[r1 - 1:]

    # mesclas (do Excel + extras da spec, p.ex. área de observações)
    faixas = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    for extra in spec.get('mesclas_extras', []):
        a, b, c, d = range_boundaries(extra)
        faixas.append((b, a, d, c))
    mescla, coberta = {}, set()
    for (mr1, mc1, mr2, mc2) in faixas:
        if mr1 > r2 or mc1 > c2 or mr2 < r1 or mc2 < c1:
            continue
        mr2, mc2 = min(mr2, r2), min(mc2, c2)
        mescla[(mr1, mc1)] = (mr2, mc2)
        for rr in range(mr1, mr2 + 1):
            for cc in range(mc1, mc2 + 1):
                if (rr, cc) != (mr1, mc1):
                    coberta.add((rr, cc))

    # caixas de seleção do Excel → papel "verificacao" (automático; a spec pode listar ou desligar com false)
    ver = spec.get('verificacoes', 'auto')
    exemplo_ver = {}
    if ver == 'auto':
        # caixas em células que a spec já definiu (p.ex. Sim/Não como "escolhas" na FR-IMOB-13) ficam como estão
        ja_definidas = papeis_da_spec(spec)
        achadas = caixas_de_selecao(caminho, ws)
        ver = []
        for a, marcada in achadas:
            cc, rr = separar(a)
            for (mr1, mc1), (mr2, mc2) in mescla.items():   # dentro de uma mescla → a célula-âncora
                if mr1 <= rr <= mr2 and mc1 <= cc <= mc2:
                    a = endereco(mc1, mr1)
                    break
            if a in ja_definidas or a in ver:
                continue
            ver.append(a)
            if marcada:
                exemplo_ver[a] = True
    spec = {**spec, '_verificacoes': ver or []}

    papeis = papeis_da_spec(spec)
    limpar = set(spec.get('limpar', []))
    formulas_spec = spec.get('formulas', {})   # {"N20": "=IF(W20=\"\",\"\",W20/0.07854*0.09807)"}
    cells, exemplo, excel, alertas = {}, {}, {}, []

    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            if (r, c) in coberta:
                continue
            cel = ws.cell(r, c)
            a = endereco(c, r)
            d = {}
            mr = mescla.get((r, c))
            if mr:
                if mr[0] > r: d['rs'] = mr[0] - r + 1
                if mr[1] > c: d['cs'] = mr[1] - c + 1
            f = fonte(cores, cel.font)
            if f: d['f'] = f
            if cel.fill is not None and cel.fill.fill_type == 'solid':
                bg = cores(cel.fill.fgColor)
                if bg and bg.upper() != '#FFFFFF': d['bg'] = bg
            ult = ws.cell(mr[0], c) if mr else cel
            dir_ = ws.cell(r, mr[1]) if mr else cel
            bd = {}
            for k, src, lado in (('t', cel, 'top'), ('l', cel, 'left'), ('b', ult, 'bottom'), ('r', dir_, 'right')):
                v = borda(cores, getattr(src.border, lado))
                if v: bd[k] = v
            if bd: d['bd'] = bd
            al = cel.alignment
            if al is not None:
                if al.horizontal and al.horizontal != 'general': d['h'] = al.horizontal
                if al.vertical and al.vertical != 'bottom': d['vt'] = al.vertical
                if al.wrap_text: d['w'] = 1
                if al.text_rotation: d['rot'] = al.text_rotation
                if al.indent: d['ind'] = al.indent
            if cel.number_format and cel.number_format != 'General':
                d['nf'] = cel.number_format

            v = cel.value
            papel = papeis.get(a)
            if a in limpar:
                v = None
            if a in formulas_spec:
                v = formulas_spec[a]
                d['fxi'] = 1   # fórmula acrescentada pela spec (não existe no Excel original)
            if isinstance(v, str) and v.startswith('='):
                if papel and papel['tipo'] in ('entrada', 'revisao'):
                    alertas.append(f'{a}: célula de entrada tinha fórmula no Excel ({v}); a fórmula foi descartada.')
                    v = None
                elif any(x in v.upper() for x in FUNCOES_BLOQUEADAS):
                    alertas.append(f'{a}: fórmula bloqueada (aleatória/volátil/dinâmica): {v}')
                    v = None
                else:
                    d['fx'] = v[1:]
                    if not d.get('fxi'):
                        cv = wsv.cell(r, c).value
                        excel[a] = serial_excel(cv) if cv is not None else None
                        if isinstance(cv, str) and cv.startswith('#'):
                            alertas.append(f'{a}: o Excel mostra {cv} nesta fórmula.')
            elif papel and papel.get('_manter'):
                # rótulo e valor na mesma célula ("Nº O.S: ____"): o texto vira prefixo
                papel = dict(papel)
                papel.pop('_manter')
                papel['prefixo'] = str(v).rstrip() if v is not None else ''
                papeis[a] = papel
                v = None
            elif papel and papel['tipo'] == 'verificacao':
                # texto na própria célula da caixa (raro) vira o rótulo ao lado do ☐
                if v is not None and str(v).strip():
                    papel = {**papel, 'texto': str(v).strip()}
                    papeis[a] = papel
                v = None
            elif papel and papel['tipo'] in ('entrada', 'pedido', 'revisao', 'escolha'):
                if v is not None:
                    exemplo[a] = str(v) if isinstance(v, CellRichText) else serial_excel(v)
                v = None if papel['tipo'] != 'escolha' or papel.get('marca') else papel.get('texto')
            if isinstance(v, str) and '#REF!' in v:
                alertas.append(f'{a}: contém #REF! ({v})')
            if isinstance(v, CellRichText):
                d['rt'] = trechos(v, cores)
            elif v is not None and 'fx' not in d:
                d['v'] = serial_excel(v)
            if papel:
                d['role'] = {k: x for k, x in papel.items() if not k.startswith('_')}
            if d or mr:
                cells[a] = d

    # rótulos das entradas (para a visão em lista no celular)
    #   rl = rótulo da linha (texto mais próximo à esquerda; se for curto, como "I" ou "Média",
    #        junta o anterior: "Espessura (mm) · I")
    #   rc = cabeçalho da coluna (texto acima, em célula estreita — faixas largas são ignoradas)
    rotulos_spec = spec.get('rotulos', {})
    def texto_em(rr, cc):
        ar, ac, span = rr, cc, 1
        for (mr1_, mc1_), (mr2_, mc2_) in mescla.items():
            if mr1_ <= rr <= mr2_ and mc1_ <= cc <= mc2_:
                ar, ac, span = mr1_, mc1_, mc2_ - mc1_ + 1
                break
        d = cells.get(endereco(ac, ar))
        if not d or d.get('role') or d.get('fx'):
            return None, span, (ar, ac)
        if 'rt' in d:
            t = ''.join(x['t'] for x in d['rt']).strip()
        else:
            v = d.get('v')
            t = str(v).strip() if isinstance(v, (str, int)) and not isinstance(v, bool) else ''
        t = re.sub(r'\s+', ' ', t).rstrip(':').strip()
        return (t or None), span, (ar, ac)
    for a, d in cells.items():
        papel = d.get('role')
        if not papel or papel['tipo'] not in ('entrada', 'revisao'):
            continue
        c, r = separar(a)
        textos, vistos = [], set()
        for cc in range(c - 1, c1 - 1, -1):
            t, _, anc = texto_em(r, cc)
            if t and anc not in vistos:
                vistos.add(anc)
                textos.append(t)
                if len(textos) == 2 or len(textos[0]) > 5:
                    break
        rl = ' · '.join(reversed(textos[:2] if textos and len(textos[0]) <= 5 else textos[:1]))
        rc = None
        for rr in range(r - 1, max(r1, r - 6) - 1, -1):
            t, span, _ = texto_em(rr, c)
            if t and span <= 3 and re.search(r'[A-Za-zÀ-ú]', t):
                rc = t
                break
        if rl: papel['rl'] = rl
        if rc and rc != rl: papel['rc'] = rc
        if a in rotulos_spec:
            papel['rot'] = rotulos_spec[a]

    # imagens (descarta as que caem nas linhas de assinatura: são assinaturas de exemplo)
    x_off = [0]
    for w in colpx: x_off.append(x_off[-1] + w)
    y_off = [0]
    for h in rowpx: y_off.append(y_off[-1] + h)
    faixas_ass = spec.get('linhas_assinatura', [])
    imgs = []
    for im in ws._images:
        an = im.anchor
        fr = an._from
        if any(lo <= fr.row + 1 <= hi for lo, hi in faixas_ass):
            continue
        clamp_c = lambda i: max(0, min(i, len(colpx)))
        clamp_r = lambda i: max(0, min(i, len(rowpx)))
        x = x_off[clamp_c(fr.col - (c1 - 1))] + fr.colOff / EMU_PX
        y = y_off[clamp_r(fr.row - (r1 - 1))] + fr.rowOff / EMU_PX
        if getattr(an, 'to', None) is not None:
            to = an.to
            w = x_off[clamp_c(to.col - (c1 - 1))] + to.colOff / EMU_PX - x
            h = y_off[clamp_r(to.row - (r1 - 1))] + to.rowOff / EMU_PX - y
        else:
            w, h = an.ext.width / EMU_PX, an.ext.height / EMU_PX
        dados = im._data()
        mime = 'image/png' if dados[:4] == b'\x89PNG' else 'image/jpeg'
        imgs.append({'x': round(x, 1), 'y': round(y, 1), 'w': round(w, 1), 'h': round(h, 1),
                     'src': f'data:{mime};base64,' + base64.b64encode(dados).decode()})

    pm = ws.page_margins
    ps = ws.page_setup
    folha = {
        'origem': {'c1': c1, 'r1': r1}, 'cols': colpx, 'rows': rowpx, 'cells': cells, 'imgs': imgs,
        **({'impressao': {'cols': c2_impressao - c1 + 1}} if c2 > c2_impressao else {}),
        'pagina': {
            'orient': ps.orientation or 'portrait',
            'margens': [pm.top, pm.right, pm.bottom, pm.left],
            'centralizar': bool(ws.print_options.horizontalCentered),
            'ajuste': [1 if ps.fitToWidth is None else ps.fitToWidth, 1 if ps.fitToHeight is None else ps.fitToHeight],
        },
        'lista': spec.get('lista', {}),
    }
    for a in spec.get('pedido', {}):
        if a not in cells:
            alertas.append(f'{a}: célula de pedido fora da área de impressão ou coberta por mescla.')
    for a, p in papeis.items():
        if a not in cells:
            alertas.append(f'{a}: célula com papel "{p["tipo"]}" fora da área de impressão ou coberta por mescla.')

    exemplo.update({'__verificacoes__': exemplo_ver} if exemplo_ver else {})
    return folha, exemplo, excel, alertas


CHAVES_DA_FOLHA = ('pedido', 'entradas', 'revisao', 'escolhas', 'assinaturas', 'linhas_assinatura', 'formulas',
                   'colunas_tela', 'limpar', 'mesclas_extras', 'lista', 'rotulos', 'verificacoes', 'area_impressao')
RE_ABA_FORMULA = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!\$?[A-Z]{1,3}\$?\d")


def converter(spec, arquivo_spec):
    """Ficha completa: aba principal (spec) + abas extras (spec['abas_extras'], p.ex. o verso)."""
    caminho = AQUI / 'planilhas' / spec['arquivo']
    wb = openpyxl.load_workbook(caminho, rich_text=True)
    wbv = openpyxl.load_workbook(caminho, data_only=True)
    cores = Cores(wb)
    ws = wb[spec['aba']] if spec.get('aba') else wb.worksheets[0]
    folha, exemplo, excel, alertas = converter_folha(caminho, wb, wbv, ws, spec, cores)
    modelo = {'motor': MOTOR, 'codigo': spec['codigo'], 'versao': spec['versao'], 'titulo': spec['nome'], **folha}
    verif_ex = exemplo.pop('__verificacoes__', {})

    extras = spec.get('abas_extras', [])
    apelidos = {ws.title: ''}
    if extras:
        modelo['titulo_aba'] = spec.get('titulo_aba', 'Frente')
        modelo['abas'] = []
        for ex in extras:
            ident = ex['id']
            if not re.fullmatch(r'[A-Z][A-Z0-9_]*', ident):
                raise SystemExit(f'{spec["codigo"]}: id de aba inválido "{ident}" (use letras maiúsculas, p.ex. "VERSO").')
            wsx = wb[ex['aba']]
            apelidos[wsx.title] = ident
            sub = {k: ex[k] for k in CHAVES_DA_FOLHA if k in ex}
            fx, exx, xlx, alx = converter_folha(caminho, wb, wbv, wsx, sub, cores)
            modelo['abas'].append({'id': ident, 'titulo': ex.get('titulo', ident), **fx})
            ver_x = exx.pop('__verificacoes__', {})
            exemplo.update({f'{ident}!{a}': v for a, v in exx.items()})
            excel.update({f'{ident}!{a}': v for a, v in xlx.items()})
            verif_ex.update({f'{ident}!{a}': v for a, v in ver_x.items()})
            alertas += [f'[{ex.get("titulo", ident)}] {m}' for m in alx]
        modelo['apelidos'] = apelidos

    # fórmulas que apontam para abas que não fazem parte da ficha
    todas = [('', modelo['cells'])] + [(ab['titulo'] + ' ', ab['cells']) for ab in modelo.get('abas', [])]
    for rot, cells in todas:
        for a, d in cells.items():
            for m in RE_ABA_FORMULA.finditer(d.get('fx', '')):
                nome = (m.group(1) or '').replace("''", "'") or m.group(2)
                if nome not in apelidos:
                    alertas.append(f'{rot}{a}: fórmula aponta para a aba "{nome}", que não faz parte da ficha (=' + d['fx'] + ')')

    verificacao = {'exemplo': exemplo, 'excel': excel, 'pedido_exemplo': spec.get('pedido_exemplo', {}),
                   **({'verificacoes': verif_ex} if verif_ex else {}),
                   'abas_excel': {v: k for k, v in apelidos.items()}}
    return modelo, spec.get('resultados', []), verificacao, alertas


# ── SQL de carga ─────────────────────────────────────────────────────────────

def sql_literal(s):
    return "'" + s.replace("'", "''") + "'"


def gerar_sql(itens):
    linhas = [
        '-- =============================================================================',
        '-- CNRO Lab Control — Migração 13b: carga dos modelos das fichas online',
        '-- =============================================================================',
        '-- GERADO AUTOMATICAMENTE por tools/fichas/converter.py — não edite à mão.',
        f'-- Gerado em {datetime.datetime.now().strftime("%d/%m/%Y %H:%M")} · motor {MOTOR}',
        '-- Pré-requisito: migração 13. Idempotente: publicar de novo a mesma revisão',
        '-- atualiza o modelo; uma revisão nova (Rev01) vira um modelo novo e passa a',
        '-- valer para os ensaios iniciados depois dela.',
        '-- Fichas: ' + ', '.join(f'{m["codigo"]} {m["versao"]}' for m, _, _ in itens),
        '-- =============================================================================',
        '',
        'begin;',
        '',
    ]
    for modelo, mapa, spec in itens:
        js = json.dumps(modelo, ensure_ascii=False, separators=(',', ':'))
        mp = json.dumps(mapa, ensure_ascii=False, separators=(',', ':'))
        h = hashlib.sha256((js + mp).encode()).hexdigest()[:16]
        cod, ver = modelo['codigo'], modelo['versao']
        linhas += [
            f'-- {cod} {ver} — {modelo["titulo"]}',
            f'insert into public.fichas_ensaio (codigo, nome, versao)',
            f'values ({sql_literal(cod)}, {sql_literal(spec.get("nome_cadastro", modelo["titulo"]))}, {sql_literal(ver)})',
            f'on conflict (codigo) do nothing;',
            '',
            'insert into public.fichas_modelo (ficha_ensaio_id, codigo, versao, titulo, modelo, mapa_resultados, motor, hash)',
            f'select f.id, {sql_literal(cod)}, {sql_literal(ver)}, {sql_literal(modelo["titulo"])},',
            f'       {sql_literal(js)}::jsonb,',
            f'       {sql_literal(mp)}::jsonb, {MOTOR}, {sql_literal(h)}',
            f'  from public.fichas_ensaio f where f.codigo = {sql_literal(cod)}',
            'on conflict (ficha_ensaio_id, versao) do update',
            '   set modelo = excluded.modelo, mapa_resultados = excluded.mapa_resultados, titulo = excluded.titulo,',
            '       motor = excluded.motor, hash = excluded.hash, ativo = true',
            ' where public.fichas_modelo.hash is distinct from excluded.hash;',
            '',
            f'update public.fichas_ensaio set versao = {sql_literal(ver)}, updated_at = now()',
            f' where codigo = {sql_literal(cod)} and versao is distinct from {sql_literal(ver)};',
            '',
        ]
    linhas += ['commit;', '', '-- Fim da migração 13b.', '']
    return '\n'.join(linhas)


def main():
    filtro = set(sys.argv[1:])
    itens, problemas = [], 0
    (AQUI / 'saida').mkdir(exist_ok=True)
    for arq in sorted((AQUI / 'specs').glob('*.json')):
        spec = json.loads(arq.read_text(encoding='utf-8'))
        if filtro and spec['codigo'] not in filtro:
            continue
        modelo, mapa, verif, alertas = converter(spec, arq)
        nome = f'{spec["codigo"]}_{spec["versao"]}'
        (AQUI / 'saida' / f'{nome}.modelo.json').write_text(
            json.dumps({'modelo': modelo, 'mapa_resultados': mapa}, ensure_ascii=False), encoding='utf-8')
        (AQUI / 'saida' / f'{nome}.verificacao.json').write_text(json.dumps(verif, ensure_ascii=False), encoding='utf-8')
        n_fx = sum(1 for c in modelo['cells'].values() if 'fx' in c)
        n_in = sum(1 for c in modelo['cells'].values() if c.get('role', {}).get('tipo') == 'entrada')
        for ab in modelo.get('abas', []):
            n_fx += sum(1 for c in ab['cells'].values() if 'fx' in c)
            n_in += sum(1 for c in ab['cells'].values() if c.get('role', {}).get('tipo') == 'entrada')
        n_ver = sum(1 for f in [modelo] + modelo.get('abas', []) for c in f['cells'].values()
                    if c.get('role', {}).get('tipo') == 'verificacao')
        abas = f' · abas: {modelo.get("titulo_aba", "Frente")} + ' + ', '.join(a['titulo'] for a in modelo['abas']) if modelo.get('abas') else ''
        print(f'{nome}: {len(modelo["cells"]) + sum(len(a["cells"]) for a in modelo.get("abas", []))} células · {n_fx} fórmulas · '
              f'{n_in} entradas{f" · {n_ver} verificações" if n_ver else ""} · '
              f'{len(modelo["imgs"])} imagem(ns) · {len(mapa)} mapa(s) de resultado{abas}')
        for a in alertas:
            print('   ⚠', a)
        problemas += len(alertas)
        itens.append((modelo, mapa, spec))
    if not filtro:
        (RAIZ / 'supabase' / 'migrations' / '13b_fichas_modelo_carga.sql').write_text(gerar_sql(itens), encoding='utf-8')
        print('SQL de carga: supabase/migrations/13b_fichas_modelo_carga.sql')


if __name__ == '__main__':
    main()
