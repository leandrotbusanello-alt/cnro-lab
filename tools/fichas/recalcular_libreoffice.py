#!/usr/bin/env python3
"""
Conferência independente do motor: grava valores de teste na planilha original,
recalcula no LibreOffice (headless) e devolve o valor de cada fórmula.

Uso (chamado por testar_motor.mjs --libreoffice):
    python3 recalcular_libreoffice.py <planilha.xlsx> <aba|""> <entradas.json> <saida.json>
entradas.json = { "D18": 65.13, "O7": 46094, "F10": null, "FR-IMOB-06 VERSO!F7": "BAL-01", ... }
    (datas como número serial do Excel; null apaga a célula; "Aba!A1" grava em outra aba)
Saída: fórmulas da aba principal pelo endereço ("D22"); das outras abas, "Aba!A1".
"""
import datetime, json, shutil, subprocess, sys, tempfile
from pathlib import Path
import openpyxl


def main():
    xlsx, aba, arq_in, arq_out = sys.argv[1:5]
    entradas = json.loads(Path(arq_in).read_text(encoding='utf-8'))
    with tempfile.TemporaryDirectory() as tmp:
        origem = Path(tmp) / 'entrada.xlsx'
        wb = openpyxl.load_workbook(xlsx)
        ws = wb[aba] if aba else wb.worksheets[0]
        for a, v in entradas.items():
            if '!' in a:
                nome, cel = a.rsplit('!', 1)
                wb[nome][cel].value = v
            else:
                ws[a].value = v
        wb.save(origem)
        saida_dir = Path(tmp) / 'out'
        saida_dir.mkdir()
        subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(saida_dir), str(origem)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        wbv = openpyxl.load_workbook(saida_dir / 'entrada.xlsx', data_only=True)
        wbf = openpyxl.load_workbook(origem)
        out = {}
        for i, wsf in enumerate(wbf.worksheets):
            wsv = wbv.worksheets[i]
            pre = '' if i == wb.worksheets.index(ws) else f'{wsf.title}!'
            for row in wsf.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith('='):
                        v = wsv[c.coordinate].value
                        if hasattr(v, 'year'):
                            base = datetime.datetime(1899, 12, 30)
                            dt = v if isinstance(v, datetime.datetime) else datetime.datetime(v.year, v.month, v.day)
                            v = (dt - base).days + (dt - base).seconds / 86400
                            if v < 61:          # Excel conta 29/02/1900 (dia inexistente): seriais < 61 ficam 1 a menos
                                v -= 1
                        out[pre + c.coordinate] = v
        Path(arq_out).write_text(json.dumps(out, ensure_ascii=False), encoding='utf-8')


if __name__ == '__main__':
    main()
