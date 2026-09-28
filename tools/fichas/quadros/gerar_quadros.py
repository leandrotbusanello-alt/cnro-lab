#!/usr/bin/env python3
"""
Quadros de controle (FR-IMOB-39 a 43): planilha da Qualidade → modelo de layout para a tela
"Quadros de controle" do Laboratório (src/modules/quadros).

  python3 tools/fichas/quadros/gerar_quadros.py            # todos
  python3 tools/fichas/quadros/gerar_quadros.py FR-IMOB-43 # um

O quadro não é digitado: o sistema preenche as células com os resultados aprovados. Aqui só se lê o
desenho (conversor de fichas com `so_layout`) e se marcam as células que recebem valor:
  - linhas de amostra (bloco de 1 linha; na FR-43, 2 linhas por série) — compactadas para UM bloco
    modelo; a tela repete o bloco quantas vezes precisar (motorQuadros.expandirModelo);
  - estatísticas (FR-39 a 42: linhas "Média Período" e "Análise Global"; a função de cada célula —
    média, máx., mín., desvio — é lida da fórmula da planilha, só para saber o que calcular);
  - cabeçalho (FR-43: traço, lançamento, estatística ACI 214).
Saída: src/modules/quadros/modelos/<CODIGO>.json = { modelo, quadro: { bloco, colunas, estatisticas, … } }.
"""
import json, re, sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[2]
sys.path.insert(0, str(AQUI.parent))
import converter as conv  # noqa: E402

PLANILHAS = AQUI / 'planilhas'
SAIDA = RAIZ / 'src' / 'modules' / 'quadros' / 'modelos'

FUNCOES = {'AVERAGE': 'media', 'MAX': 'max', 'MIN': 'min', 'STDEV': 'desvio'}

# ── Definição do layout de cada quadro ───────────────────────────────────────
#   bloco: primeira linha de amostra, linhas por amostra, última linha de amostra na planilha
#   colunas: tipo de cada coluna do bloco (numero | texto | data)
#   estat: linhas de estatística → 'periodo' | 'global'
QUADROS = {
    'FR-IMOB-39': dict(arquivo='FR-IMOB-39_Rev00.xlsx', aba='FR-IMOB-39_Rev00', nome='Caracterização do agregado',
                       bloco=(13, 1, 172), colunas={**{c: 'numero' for c in 'CEFGHIJKLMNOPQRS'}, 'B': 'data', 'D': 'data', 'T': 'texto', 'U': 'texto'},
                       estat={173: 'periodo', 174: 'global', 175: 'global', 176: 'global', 177: 'global'}),
    'FR-IMOB-40': dict(arquivo='FR-IMOB-40_Rev00.xlsx', aba='FR-IMOB-40_REV00', nome='Caracterização do CAP',
                       bloco=(13, 1, 172), colunas={**{c: 'numero' for c in 'CEFGHIJKL'}, 'B': 'data', 'D': 'data', 'M': 'texto', 'N': 'texto'},
                       estat={173: 'periodo', 174: 'global', 175: 'global', 176: 'global', 177: 'global'}),
    'FR-IMOB-41': dict(arquivo='FR-IMOB-41_Rev00.xlsx', aba='FR-IMOB-41_Rev00', nome='Caracterização do CBUQ endurecido',
                       bloco=(13, 1, 172), colunas={**{c: 'numero' for c in 'BDEFGHIJKLMNOPQRSTUVW'}, 'A': 'data', 'C': 'data', 'X': 'texto', 'Y': 'texto'},
                       estat={173: 'periodo', 174: 'global', 175: 'global', 176: 'global', 177: 'global'}),
    'FR-IMOB-42': dict(arquivo='FR-IMOB-42_Rev00.xlsx', aba='FR-IMOB-42_Rev00', nome='Caracterização do CBUQ fresco',
                       bloco=(13, 1, 172), colunas={**{c: 'numero' for c in 'CEFGHIJKLMNOPQRSTUVWX'}, 'B': 'data', 'D': 'data', 'Y': 'texto', 'Z': 'texto'},
                       estat={173: 'periodo', 174: 'global', 175: 'global', 176: 'global', 177: 'global'}),
    # FR-43: série = 2 linhas (CP 1 e CP 2). Colunas mescladas nas 2 linhas ficam só na 1ª.
    'FR-IMOB-43': dict(arquivo='FR-IMOB-43_Rev00.xlsx', aba=None, nome='Controle de qualidade do concreto e argamassa',
                       bloco=(18, 2, 77),
                       colunas={'A18': 'numero', 'B18': 'data', 'C18': 'texto', 'D18': 'numero', 'E18': 'numero', 'F18': 'numero',
                                'G18': 'texto', 'G19': 'texto', 'I18': 'numero',
                                'L18': 'numero', 'L19': 'numero', 'M18': 'numero', 'N18': 'numero', 'N19': 'numero', 'O18': 'numero',
                                'P18': 'numero', 'P19': 'numero', 'Q18': 'numero', 'R18': 'numero', 'S18': 'texto', 'T18': 'texto',
                                'U18': 'texto', 'U19': 'texto', 'V18': 'texto'},
                       cabecalho={'C9': 'texto', 'G9': 'texto', 'S9': 'texto', 'U9': 'texto',
                                  'E12': 'texto', 'E13': 'texto', 'E14': 'texto',
                                  'I11': 'numero', 'I12': 'numero', 'I13': 'numero', 'I14': 'numero',
                                  'L11': 'numero', 'L12': 'numero', 'L13': 'numero', 'L14': 'numero',
                                  'N11': 'numero', 'N12': 'numero', 'N13': 'numero', 'N14': 'numero',
                                  'P11': 'numero', 'P12': 'numero', 'P13': 'numero', 'P14': 'numero', 'P15': 'numero',
                                  'S11': 'texto', 'S12': 'numero', 'S13': 'numero', 'S14': 'numero', 'T13': 'numero',
                                  'W13': 'numero', 'X13': 'texto', 'C15': 'numero'},
                       limpar=['U15'],      # "CQ - Holanda Engenharia Ltda" (texto de outra empresa na planilha)
                       # ajustes de layout (textos que no Excel transbordam para a célula vazia ao lado; o rótulo do
                       # consumo de cimento transbordava para a esquerda: vai para N9:R9, alinhado à direita)
                       ajustes={'C9': {'f.s': 11, 'w': 1}, 'N9': {'copiar': 'R9', 'cs': 5}, 'R9': {'apagar': True}, 'A15': {'copiar': 'B15', 'cs': 2, 'bd.l': 'A15'}, 'B15': {'apagar': True},
                                'E12': {'cs': 2, 'bd.r': 'F12'}, 'E13': {'cs': 2, 'bd.r': 'F13'}, 'E14': {'cs': 2, 'bd.r': 'F14'}},
                       formatos={'S12': '0', 'S13': '0.0', 'S14': '0.0', 'T13': '0.0', 'P15': '0.0', 'I14': '0.0', 'L14': '0.0', 'N14': '0.0', 'P14': '0.0',
                                 'I12': '0.0', 'I13': '0.0', 'L12': '0.0', 'L13': '0.0', 'N12': '0.0', 'N13': '0.0', 'P12': '0.0', 'P13': '0.0',
                                 'W13': '0', 'D18': 'h:mm'}),
}


def ler_estatisticas(ws, estat, colunas_bloco):
    """{ 'E173': {'fn': 'media', 'escopo': 'periodo', 'col': 'E'} } a partir das fórmulas da planilha."""
    out = {}
    for r, escopo in estat.items():
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            t = getattr(v, 'text', v)
            if not (isinstance(t, str) and t.startswith('=')):
                continue
            m = re.match(r'=\s*(AVERAGE|MAX|MIN|STDEV)\(\s*([A-Z]+)\d+:', t.upper())
            col = cl(c)
            if m and col in colunas_bloco:
                out[f'{col}{r}'] = {'fn': FUNCOES[m.group(1)], 'escopo': escopo, 'col': col}
    return out


def compactar(modelo, linha, altura, fim, meta):
    """Deixa só o 1º bloco de amostra: tira as linhas do 2º bloco até `fim` e sobe o que vem abaixo."""
    r1 = modelo['origem']['r1']
    ini_rem, fim_rem = linha + altura, fim            # linhas removidas (1-based, inclusive)
    n = fim_rem - ini_rem + 1
    if n <= 0:
        return modelo
    px_rem = sum(modelo['rows'][ini_rem - r1: fim_rem - r1 + 1])
    novas = {}
    for a, d in modelo['cells'].items():
        col, r = re.match(r'([A-Z]+)(\d+)', a).groups()
        r = int(r)
        if ini_rem <= r <= fim_rem:
            continue
        rs = d.get('rs', 1)
        if r < ini_rem and r + rs - 1 >= ini_rem:            # mescla que atravessa a faixa removida
            d = {**d, 'rs': max(1, rs - min(n, r + rs - ini_rem))}
        novas[f'{col}{r - n if r > fim_rem else r}'] = d
    modelo['cells'] = novas
    del modelo['rows'][ini_rem - r1: fim_rem - r1 + 1]
    topo = sum(modelo['rows'][: ini_rem - r1])
    for im in modelo.get('imgs', []):
        if im['y'] >= topo + px_rem:
            im['y'] -= px_rem
    for g in modelo.get('graficos', []):
        if g['y'] >= topo + px_rem:
            g['y'] -= px_rem
    # estatísticas e demais endereços abaixo do bloco sobem n linhas
    meta['estatisticas'] = {
        (f'{k[:len(k.rstrip("0123456789"))]}{int(k[len(k.rstrip("0123456789")):]) - n}'
         if int(k[len(k.rstrip("0123456789")):]) > fim_rem else k): v
        for k, v in meta.get('estatisticas', {}).items()}
    return modelo


def gerar(codigo, q):
    caminho = PLANILHAS / q['arquivo']
    wb = openpyxl.load_workbook(caminho)
    ws = wb[q['aba']] if q['aba'] else wb.worksheets[0]
    linha, altura, fim = q['bloco']

    # células que recebem valor (no bloco modelo e fora dele)
    if all(re.fullmatch(r'[A-Z]+', k) for k in q['colunas']):
        bloco = {f'{c}{linha}': t for c, t in q['colunas'].items()}
    else:
        bloco = dict(q['colunas'])
    estat = ler_estatisticas(ws, q.get('estat', {}), {re.match(r'[A-Z]+', k).group() for k in bloco})
    cab = dict(q.get('cabecalho', {}))
    papeis = {**bloco, **{a: 'numero' for a in estat}, **cab}

    por_tipo = {}
    for a, t in papeis.items():
        por_tipo.setdefault(t, []).append(a)
    spec = {
        # o conversor lê de tools/fichas/planilhas: caminho relativo até tools/fichas/quadros/planilhas
        'codigo': codigo, 'versao': 'Rev00', 'nome': q['nome'], 'arquivo': f"../quadros/planilhas/{q['arquivo']}",
        **({'aba': q['aba']} if q['aba'] else {}),
        'so_layout': True, 'verificacoes': False, 'graficos': False,
        'entradas': [{'celulas': sorted(v), 'tipo': t} for t, v in por_tipo.items()],
        'limpar': q.get('limpar', []), 'formatos': q.get('formatos', {}),
    }
    modelo, _, _, alertas = conv.converter(spec, None)
    modelo['pagina']['ajuste'] = [1, 0]           # largura da página; altura cresce com as amostras (quebra de página)
    meta = {'bloco': {'linha': linha, 'altura': altura, 'capacidade_planilha': (fim - linha + 1) // altura},
            'celulas_bloco': bloco, 'cabecalho': cab, 'estatisticas': estat}
    for a, aj in q.get('ajustes', {}).items():
        if aj.get('apagar'):
            modelo['cells'].pop(a, None)
            continue
        if aj.get('copiar'):
            esq = modelo['cells'].get(a, {}).get('bd', {}).get('l')
            modelo['cells'][a] = json.loads(json.dumps(modelo['cells'][aj['copiar']]))
            if esq:
                modelo['cells'][a].setdefault('bd', {})['l'] = esq
        d = modelo['cells'].setdefault(a, {})
        for k, v in aj.items():
            if k == 'copiar':
                continue
            if k == 'f.s':
                d.setdefault('f', {})['s'] = v
            elif k == 'bd.l':
                continue        # tratado no 'copiar'
            elif k == 'bd.r':
                viz = modelo['cells'].get(v, {}).get('bd', {}).get('r')
                if viz:
                    d.setdefault('bd', {})['r'] = viz
            else:
                d[k] = v
    compactar(modelo, linha, altura, fim, meta)
    # entradas sem rótulos de lista (a tela não usa a visão em lista)
    for d in modelo['cells'].values():
        if d.get('role', {}).get('tipo') == 'entrada':
            for k in ('rl', 'rc'):
                d['role'].pop(k, None)
    SAIDA.mkdir(parents=True, exist_ok=True)
    arq = SAIDA / f'{codigo}.json'
    arq.write_text(json.dumps({'modelo': modelo, 'quadro': meta}, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'{codigo}: {len(modelo["cells"])} células · bloco {linha}+{altura} · {len(estat)} estatísticas · '
          f'{len(cab)} de cabeçalho · {arq.stat().st_size // 1024} KB')
    for a in alertas:
        if 'fórmula' not in a:
            print('   ⚠', a)


if __name__ == '__main__':
    filtro = set(sys.argv[1:])
    for codigo, q in QUADROS.items():
        if not filtro or codigo in filtro:
            gerar(codigo, q)
